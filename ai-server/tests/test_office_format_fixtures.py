"""R0-5: office(docx/pptx/xlsx) 형식 표 추출 스모크.

배경
----
RAG 품질 개선이 PDF(모집요강) 중심으로 누적되면서(예: #103 OpenDataLoader JSON
색인은 PDF 전용), office 문서 경로가 조용히 망가져도 평가셋(전부 PDF)으로는
드러나지 않는다. office 문서는 학교 챗봇의 핵심 컨텐츠(장학/취업 통계 xlsx,
학사안내 docx, 학과소개 pptx)이므로 그 경로를 최소한으로라도 잠가 둔다.

이 테스트가 잠그는 것
---------------------
각 office 형식에 대해 공유 파이프라인의 앞단을 end-to-end 로 확인한다.
1. ``_load_documents`` 가 office 파일을 Markdown 으로 변환한다(MarkItDown).
2. 표 셀 값이 변환 텍스트까지 보존된다(고유 토큰 추적).
3. 공유 경로 ``_extract_table_facts`` 가 그 표에서 table fact 를 뽑고,
   표 값이 fact 까지 흐른다.

픽스처는 런타임에 생성한다. ``Assets/`` 는 .gitignore 대상이라 실제 학교 문서를
커밋할 수 없고, 최소 합성 픽스처가 더 결정론적이며 자체완결적이다. 각 형식의
생성 라이브러리(openpyxl/python-docx/python-pptx)나 markitdown 이 없으면 해당
케이스를 명시적 사유와 함께 skip 한다 — office 업로드가 가능한 환경(운영/데스크탑)
에서는 항상 실행된다.

검증 기준은 실제 동작을 먼저 관찰(probe)해 맞췄다: 2026-06 기준 MarkItDown 은
xlsx/docx/pptx 표를 모두 Markdown 파이프 표로 변환하고, ``_extract_table_facts``
가 세 형식 모두에서 1개 이상 fact 를 만든다.
"""

from __future__ import annotations

import pytest

from main import _extract_table_facts, _load_documents

# 픽스처에만 심는 고유 토큰. 표 셀 → Markdown → table fact 까지 값이 끝까지
# 흐르는지 substring 으로 추적하기 위한 것이라, 우연히 다른 곳에 나올 수 없는 값.
DISTINCT = "Z9Q7XK0"


def _make_xlsx(path: str) -> None:
    openpyxl = pytest.importorskip("openpyxl", reason="openpyxl 없으면 xlsx 픽스처 생성 불가")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "통계"
    sheet.append(["전형", "모집인원", "경쟁률"])
    sheet.append(["수시", "100", "3.5"])
    sheet.append(["정시", DISTINCT, "4.1"])
    workbook.save(path)


def _make_docx(path: str) -> None:
    docx = pytest.importorskip("docx", reason="python-docx 없으면 docx 픽스처 생성 불가")
    document = docx.Document()
    document.add_heading("학사 안내", level=1)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "항목"
    table.rows[0].cells[1].text = "값"
    table.rows[1].cells[0].text = "등록금"
    table.rows[1].cells[1].text = DISTINCT
    document.save(path)


def _make_pptx(path: str) -> None:
    pptx = pytest.importorskip("pptx", reason="python-pptx 없으면 pptx 픽스처 생성 불가")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2)).table
    table.cell(0, 0).text = "과목"
    table.cell(0, 1).text = "학점"
    table.cell(1, 0).text = "자료구조"
    table.cell(1, 1).text = DISTINCT
    presentation.save(path)


_FIXTURE_BUILDERS = {"xlsx": _make_xlsx, "docx": _make_docx, "pptx": _make_pptx}


@pytest.mark.parametrize("ext", sorted(_FIXTURE_BUILDERS))
def test_office_format_table_fact_smoke(ext: str, tmp_path) -> None:
    pytest.importorskip("markitdown", reason="markitdown 없으면 office 파싱 경로를 탈 수 없다")

    fixture_path = tmp_path / f"fixture.{ext}"
    _FIXTURE_BUILDERS[ext](str(fixture_path))

    documents = _load_documents(str(fixture_path), fixture_path.name)
    assert documents, f"{ext}: _load_documents 가 빈 결과를 반환했다"
    text = "\n".join(doc.page_content for doc in documents)

    # (1) office -> Markdown 변환에서 표 셀 값과 표 구조가 보존됐는가
    assert DISTINCT in text, f"{ext}: 표 셀 값이 변환 텍스트에서 사라졌다(파서 회귀 의심)"
    assert "|" in text, f"{ext}: Markdown 표(|)가 생성되지 않았다"

    # (2) 공유 경로 _extract_table_facts 가 표 근거를 뽑고, 값이 fact 까지 흐르는가
    facts = _extract_table_facts(text, {"source": fixture_path.name, "document_id": "r05-fixture"})
    assert facts, f"{ext}: 표 근거(table fact)가 하나도 추출되지 않았다"
    assert any(DISTINCT in fact for fact in facts), (
        f"{ext}: 표 값이 table fact 까지 흐르지 않았다(추출 경로 회귀 의심)"
    )
